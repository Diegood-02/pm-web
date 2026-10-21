"""Pruebas de los tres formatos de salida (CSV, XLSX, PDF) y del formato JSON con esquema."""

import copy
import csv
import io
from datetime import date
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook

from app.costing.evm import EvmSummary
from app.exports.project_json import build_document, validate_document
from app.exports.schedule import CSV_HEADERS, ScheduleRow, to_csv, to_xlsx
from app.exports.summary import SummaryReport, to_pdf

ROWS = [
    ScheduleRow("1", "Diseño", False, 3, date(2026, 1, 5), date(2026, 1, 7), 0, 0, True, 50),
    ScheduleRow("2", "Construcción", False, 4, date(2026, 1, 8), date(2026, 1, 14), 0, 0, True, 0),
    ScheduleRow("3", "Pruebas", False, 2, date(2026, 1, 8), date(2026, 1, 9), 3, 3, False, 0),
]


def test_csv_tiene_encabezados_en_espanol_y_acentos_con_bom() -> None:
    content = to_csv(ROWS)

    assert content.startswith(b"\xef\xbb\xbf")
    parsed = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    assert parsed[0] == list(CSV_HEADERS)
    assert parsed[1][:4] == ["1", "Diseño", "No", "3"]
    assert parsed[1][8] == "Sí"
    assert parsed[3][6] == "3"


def test_xlsx_tiene_hoja_de_cronograma_con_fechas_y_criticas_resaltadas() -> None:
    content = to_xlsx(ROWS, "Migración ERP")

    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook["Cronograma"]
    assert [cell.value for cell in sheet[1]] == list(CSV_HEADERS)
    assert sheet["B2"].value == "Diseño"
    assert sheet["E2"].value.date() == date(2026, 1, 5)
    assert sheet["E2"].number_format == "yyyy-mm-dd"
    assert sheet["A2"].fill.start_color.rgb.endswith("FDE2E2")
    assert sheet["A4"].fill.fill_type is None
    info = workbook["Proyecto"]
    assert info["B1"].value == "Migración ERP"
    assert info["B3"].value == 2


def test_xlsx_sin_tareas_sigue_siendo_valido() -> None:
    workbook = load_workbook(io.BytesIO(to_xlsx([], "Vacío")))

    assert workbook["Cronograma"].max_row == 1


def _report(evm: EvmSummary | None) -> SummaryReport:
    return SummaryReport(
        project_name="Migración ERP & Co.",
        project_start=date(2026, 1, 5),
        project_finish=date(2026, 1, 14),
        duration_days=7,
        status_date=date(2026, 1, 7),
        critical_path=["1 Diseño", "2 Construcción"],
        rows=ROWS,
        evm=evm,
    )


def test_pdf_es_un_documento_valido_con_y_sin_valor_ganado() -> None:
    evm = EvmSummary(
        bac=Decimal("4900"),
        pv=Decimal("2700"),
        ev=Decimal("2250"),
        ac=Decimal("2400"),
        cv=Decimal("-150"),
        sv=Decimal("-450"),
        cpi=0.9375,
        spi=0.8333,
        cpi_alert=False,
        spi_alert=True,
    )

    with_evm = to_pdf(_report(evm))
    without_evm = to_pdf(_report(None))

    for content in (with_evm, without_evm):
        assert content.startswith(b"%PDF-")
        assert content.rstrip().endswith(b"%%EOF")
    assert len(with_evm) > len(without_evm) * 0.8


def _document() -> dict[str, Any]:
    tasks: list[dict[str, Any]] = [
        {
            "id": 10,
            "parent_id": None,
            "name": "Fase 1",
            "duration_days": 0,
            "is_milestone": False,
            "fixed_cost": 0,
            "actual_cost": 0,
            "percent_complete": 0,
        },
        {
            "id": 11,
            "parent_id": 10,
            "name": "Diseño",
            "duration_days": 3,
            "is_milestone": False,
            "fixed_cost": 500,
            "actual_cost": 2400,
            "percent_complete": 50,
        },
        {
            "id": 12,
            "parent_id": 10,
            "name": "Construcción",
            "duration_days": 4,
            "is_milestone": False,
            "fixed_cost": 0,
            "actual_cost": 0,
            "percent_complete": 0,
        },
        {
            "id": 13,
            "parent_id": None,
            "name": "Entrega",
            "duration_days": 0,
            "is_milestone": True,
            "fixed_cost": 0,
            "actual_cost": 0,
            "percent_complete": 0,
        },
    ]
    return build_document(
        project={
            "name": "Migración ERP",
            "timezone": "America/Mexico_City",
            "start_date": "2026-01-05",
            "working_weekdays": [0, 1, 2, 3, 4],
            "hours_per_day": 8,
            "holidays": ["2026-01-06"],
        },
        tasks=tasks,
        resources_list=[{"id": 1, "name": "Ana", "max_units": 100, "hourly_rate": 100}],
        dependencies=[
            {"predecessor_id": 11, "successor_id": 12, "type": "FS", "lag_days": 0},
            {"predecessor_id": 12, "successor_id": 13, "type": "FS", "lag_days": 0},
        ],
        assignments=[{"task_id": 11, "resource_id": 1, "units_percent": 100}],
    )


def test_documento_generado_es_valido_y_usa_referencias_locales() -> None:
    doc = _document()

    assert validate_document(doc) == []
    assert doc["tasks"][1] == {
        "ref": "T2",
        "parent_ref": "T1",
        "name": "Diseño",
        "duration_days": 3,
        "is_milestone": False,
        "fixed_cost": 500.0,
        "actual_cost": 2400.0,
        "percent_complete": 50,
    }
    assert doc["dependencies"][0] == {
        "predecessor_ref": "T2",
        "successor_ref": "T3",
        "type": "FS",
        "lag_days": 0,
    }


def test_campo_desconocido_se_rechaza() -> None:
    doc = _document()
    doc["tasks"][0]["color"] = "rojo"

    errors = validate_document(doc)

    assert any("color" in error or "Additional properties" in error for error in errors)


def test_tipo_incorrecto_y_fuera_de_rango_se_rechazan() -> None:
    doc = _document()
    doc["tasks"][1]["percent_complete"] = 150
    doc["project"]["hours_per_day"] = "ocho"

    errors = validate_document(doc)

    assert any("tasks/1/percent_complete" in error for error in errors)
    assert any("project/hours_per_day" in error for error in errors)


def test_referencia_desconocida_y_ref_repetida() -> None:
    doc = _document()
    doc["dependencies"].append(
        {"predecessor_ref": "T9", "successor_ref": "T1", "type": "FS", "lag_days": 0}
    )
    doc["resources"].append(copy.deepcopy(doc["resources"][0]))

    errors = validate_document(doc)

    assert any("referencia a tarea desconocida" in error for error in errors)
    assert any("ref repetida R1" in error for error in errors)


def test_ciclo_en_dependencias_se_rechaza() -> None:
    doc = _document()
    doc["dependencies"].append(
        {"predecessor_ref": "T4", "successor_ref": "T2", "type": "FS", "lag_days": 0}
    )

    errors = validate_document(doc)

    assert any("Ciclo detectado" in error for error in errors)


def test_ciclo_en_jerarquia_se_rechaza() -> None:
    doc = _document()
    doc["tasks"][0]["parent_ref"] = "T2"

    errors = validate_document(doc)

    assert any("jerarquía circular" in error for error in errors)


def test_hito_con_duracion_y_zona_horaria_invalida() -> None:
    doc = _document()
    doc["tasks"][3]["duration_days"] = 2
    doc["project"]["timezone"] = "Marte/Olympus"

    errors = validate_document(doc)

    assert any("un hito debe tener duración 0" in error for error in errors)
    assert any("zona horaria no válida" in error for error in errors)


def test_asignacion_repetida_y_dependencia_repetida() -> None:
    doc = _document()
    doc["assignments"].append(copy.deepcopy(doc["assignments"][0]))
    doc["dependencies"].append(copy.deepcopy(doc["dependencies"][0]))

    errors = validate_document(doc)

    assert any("asignación repetida" in error for error in errors)
    assert any("dependencias repetidas" in error for error in errors)


def test_formato_o_version_incorrectos() -> None:
    doc = _document()
    doc["schema_version"] = 2
    doc["format"] = "otro"

    errors = validate_document(doc)

    assert len(errors) >= 2
    assert validate_document([]) != []
