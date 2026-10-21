"""Exportación del cronograma a CSV y a XLSX. Funciones puras."""

import csv
import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

CSV_HEADERS = (
    "WBS",
    "Tarea",
    "Hito",
    "Duración (días)",
    "Inicio",
    "Fin",
    "Holgura total (días)",
    "Holgura libre (días)",
    "Crítica",
    "Avance (%)",
)
_CRITICAL_FILL = PatternFill(start_color="FDE2E2", end_color="FDE2E2", fill_type="solid")


@dataclass(frozen=True, slots=True)
class ScheduleRow:
    wbs: str
    name: str
    is_milestone: bool
    duration_days: int
    start_date: date
    finish_date: date
    total_float: int
    free_float: int
    is_critical: bool
    percent_complete: int


def to_csv(rows: Sequence[ScheduleRow]) -> bytes:
    """CSV en UTF-8 con BOM para que Excel lo abra con los acentos correctos."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(CSV_HEADERS)
    for row in rows:
        writer.writerow(
            [
                row.wbs,
                row.name,
                "Sí" if row.is_milestone else "No",
                row.duration_days,
                row.start_date.isoformat(),
                row.finish_date.isoformat(),
                row.total_float,
                row.free_float,
                "Sí" if row.is_critical else "No",
                row.percent_complete,
            ]
        )
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(rows: Sequence[ScheduleRow], project_name: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Cronograma"
    sheet.append(list(CSV_HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for row in rows:
        sheet.append(
            [
                row.wbs,
                row.name,
                "Sí" if row.is_milestone else "No",
                row.duration_days,
                row.start_date,
                row.finish_date,
                row.total_float,
                row.free_float,
                "Sí" if row.is_critical else "No",
                row.percent_complete,
            ]
        )
        current = sheet.max_row
        for column in (5, 6):
            sheet.cell(current, column).number_format = "yyyy-mm-dd"
        if row.is_critical:
            for column in range(1, len(CSV_HEADERS) + 1):
                sheet.cell(current, column).fill = _CRITICAL_FILL

    widths = [10, 42, 8, 16, 12, 12, 20, 20, 10, 12]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"

    info = workbook.create_sheet("Proyecto")
    info.append(["Proyecto", project_name])
    info.append(["Tareas", len(rows)])
    info.append(["Tareas críticas", sum(1 for row in rows if row.is_critical)])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 50

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
